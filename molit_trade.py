"""국토교통부 아파트 매매 실거래가 조회 (공공데이터포털 RTMSDataSvcAptTradeDev API).

data.go.kr에서 발급받은 서비스키(MOLIT_API_KEY 환경변수)로 지역·월별 아파트 매매
실거래 내역을 가져온다. Playwright/브라우저가 필요 없는 순수 REST 호출이라 봇 차단
걱정이 없다 (crawl_complex.py의 네이버 크롤과 반대 — 자세한 배경은 SKILL.md 참고).

watchlist.json에 등록된 단지명과 매칭해 output/실거래가_{시각}.xlsx 리포트를 만들고,
원본 응답은 data/molit/{LAWD_CD}/{YYYYMM}.json에 누적 저장한다.
"""
import argparse
import json
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from watchlist import load_watchlist

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data" / "molit"
OUTPUT_DIR = BASE_DIR / "output"

API_URL = "http://apis.data.go.kr/1613000/RTMSDataSvcAptTradeDev/getRTMSDataSvcAptTradeDev"

# watchlist.json의 region -> 법정동코드(시군구 5자리) 매핑.
# 새 지역이 필요하면 https://www.code.go.kr (행정표준코드관리시스템)의
# "법정동코드 전체자료"에서 시군구 5자리를 찾아 여기에 추가한다.
REGION_LAWD_CD = {
    "수원광교": "41117",  # 경기도 수원시 영통구 (광교 신도시 단지 대부분 이 구 소재)
}

CELL_BORDER = Border(*(Side(style="thin"),) * 4)


def _api_request(params: dict) -> ET.Element:
    url = f"{API_URL}?{urllib.parse.urlencode(params)}"
    with urllib.request.urlopen(url, timeout=20) as resp:
        body = resp.read()
    root = ET.fromstring(body)
    result_code = root.findtext(".//resultCode")
    if result_code not in (None, "00", "000"):
        result_msg = root.findtext(".//resultMsg")
        raise RuntimeError(f"국토부 API 오류 (resultCode={result_code}): {result_msg}")
    return root


def fetch_trades(service_key: str, lawd_cd: str, year_month: str) -> list:
    """지역코드(lawd_cd, 5자리)·거래년월(year_month, YYYYMM) 기준 매매 전체를 페이지네이션해 가져온다."""
    rows = []
    page = 1
    num_of_rows = 1000
    while True:
        params = {
            "serviceKey": service_key,
            "LAWD_CD": lawd_cd,
            "DEAL_YMD": year_month,
            "pageNo": page,
            "numOfRows": num_of_rows,
        }
        root = _api_request(params)
        items = root.findall(".//item")
        for item in items:
            rows.append({tag.tag: (tag.text or "").strip() for tag in item})
        total_count = int(root.findtext(".//totalCount") or 0)
        if page * num_of_rows >= total_count or not items:
            break
        page += 1
        time.sleep(0.3)  # 공공데이터포털은 봇 차단은 없지만 과도한 연속호출은 피한다
    return rows


def simplify(raw: dict) -> dict:
    """XML item의 원본 필드를 리포트에 쓰는 형태로 정리한다."""
    deal_amount = int((raw.get("dealAmount") or "0").replace(",", "")) * 10000  # 만원 -> 원
    area = float(raw.get("excluUseAr") or 0)
    return {
        "aptName": raw.get("aptNm", ""),
        "dong": raw.get("umdNm", ""),
        "jibun": raw.get("jibun", ""),
        "aptDong": raw.get("aptDong", "").strip(),
        "floor": raw.get("floor", ""),
        "excluUseAr": area,
        "buildYear": raw.get("buildYear", ""),
        "dealDate": "{}-{:02d}-{:02d}".format(
            int(raw.get("dealYear", 0)), int(raw.get("dealMonth", 0)), int(raw.get("dealDay", 0))
        ),
        "dealAmount": deal_amount,
        "dealingGbn": raw.get("dealingGbn", "").strip() or "-",  # 중개거래/직거래
        "cancelled": raw.get("cdealType", "").strip() == "O",  # 계약 해제 여부
        "cancelDate": raw.get("cdealDay", "").strip(),
        "brokerRegion": raw.get("estateAgentSggNm", "").strip() or "-",
    }


def match_complex(apt_name: str, complex_name: str) -> bool:
    """국토부 단지명(aptNm)이 네이버 관심단지명과 같은 단지를 가리키는지 느슨하게 판단.

    국토부 공식명과 네이버 표기가 완전히 같지 않을 수 있어(예: 괄호 병기, 차수 표기
    차이) 포함 관계로 매칭한다. 오탐이 걱정되면 --lawd-cd 모드로 지역 전체를 뽑아
    직접 확인하는 편이 안전하다.
    """
    a = apt_name.replace(" ", "")
    b = complex_name.replace(" ", "")
    return a in b or b in a


def write_sheet(ws, title: str, rows: list):
    ws.title = title[:31]
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)

    headers = ["계약일", "동", "층", "전용면적(㎡)", "거래금액", "거래유형", "해제여부", "중개사소재지"]
    header_row = 3
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(bold=True)
        c.border = CELL_BORDER
        c.alignment = Alignment(horizontal="center")

    rows_sorted = sorted(rows, key=lambda r: r["dealDate"], reverse=True)
    for i, r in enumerate(rows_sorted):
        row = header_row + 1 + i
        values = [
            r["dealDate"],
            r["aptDong"] or "-",
            r["floor"],
            r["excluUseAr"],
            f"{r['dealAmount'] // 10000:,}만원",
            r["dealingGbn"],
            "해제" if r["cancelled"] else "",
            r["brokerRegion"],
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = CELL_BORDER
            if col == 5:
                cell.font = Font(bold=True, color="0000FF")
            if r["cancelled"]:
                cell.font = Font(color="FF0000")

    for col, width in enumerate([12, 10, 6, 14, 14, 10, 8, 18], start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.freeze_panes = f"A{header_row + 1}"


def build_excel(grouped: dict, out_path: Path):
    """grouped: {sheet_title: [row, ...]}"""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in grouped.items():
        ws = wb.create_sheet()
        write_sheet(ws, title, rows)
    wb.save(out_path)


def collect_year_months(months: int) -> list:
    """오늘 기준 최근 months개월(당월 포함)의 YYYYMM 리스트를 최신순으로."""
    result = []
    year, month = datetime.now().year, datetime.now().month
    for _ in range(months):
        result.append(f"{year}{month:02d}")
        month -= 1
        if month == 0:
            month = 12
            year -= 1
    return result


def main():
    parser = argparse.ArgumentParser(description="국토교통부 아파트 매매 실거래가 조회")
    parser.add_argument("--service-key", help="미지정 시 MOLIT_API_KEY 환경변수 사용")
    parser.add_argument("--months", type=int, default=1, help="오늘 기준 최근 N개월 조회 (기본 1)")
    parser.add_argument("--name", help="watchlist에 등록된 단지 이름 하나만 조회")
    parser.add_argument("--lawd-cd", help="법정동코드 5자리 직접 지정 (watchlist 무관, 지역 전체 조회)")
    parser.add_argument("--year-month", help="--lawd-cd와 함께 특정 월(YYYYMM)만 조회 (기본: --months 범위)")
    args = parser.parse_args()

    service_key = args.service_key or __import__("os").environ.get("MOLIT_API_KEY")
    if not service_key:
        print("MOLIT_API_KEY 환경변수 또는 --service-key가 필요합니다. (data.go.kr에서 발급)")
        return

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)
    year_months = [args.year_month] if args.year_month else collect_year_months(args.months)

    if args.lawd_cd:
        # 지역 전체 조회 (watchlist 매칭 없이 그 지역 모든 아파트 매매)
        all_rows = []
        for ym in year_months:
            raw_rows = fetch_trades(service_key, args.lawd_cd, ym)
            (DATA_DIR / args.lawd_cd).mkdir(exist_ok=True)
            (DATA_DIR / args.lawd_cd / f"{ym}.json").write_text(
                json.dumps(raw_rows, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            all_rows.extend(simplify(r) for r in raw_rows)
        grouped = {}
        for r in all_rows:
            grouped.setdefault(r["aptName"], []).append(r)
        if not grouped:
            print(f"{args.lawd_cd} 지역, {year_months}: 거래 없음")
            return
        out_path = OUTPUT_DIR / f"실거래가_{args.lawd_cd}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        build_excel(grouped, out_path)
        for name, rows in grouped.items():
            print(f"{name}: {len(rows)}건")
        print(f"\n완료: {out_path}")
        return

    # watchlist 연동
    watchlist = load_watchlist()
    if args.name:
        watchlist = [c for c in watchlist if c["complexName"] == args.name]
        if not watchlist:
            print(f"관심단지에 '{args.name}'이(가) 없습니다.")
            return

    regions_needed = {c.get("region", "") for c in watchlist}
    unknown_regions = regions_needed - set(REGION_LAWD_CD)
    if unknown_regions:
        print(f"REGION_LAWD_CD에 없는 지역: {unknown_regions} (molit_trade.py 상단에 추가 필요)")
        return

    lawd_cds_needed = {REGION_LAWD_CD[c["region"]] for c in watchlist}
    raw_by_lawd = {}
    for lawd_cd in lawd_cds_needed:
        combined = []
        for ym in year_months:
            raw_rows = fetch_trades(service_key, lawd_cd, ym)
            (DATA_DIR / lawd_cd).mkdir(exist_ok=True)
            (DATA_DIR / lawd_cd / f"{ym}.json").write_text(
                json.dumps(raw_rows, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            combined.extend(raw_rows)
        raw_by_lawd[lawd_cd] = combined

    grouped = {}
    for c in watchlist:
        lawd_cd = REGION_LAWD_CD[c["region"]]
        matched = [simplify(r) for r in raw_by_lawd[lawd_cd] if match_complex(r.get("aptNm", ""), c["complexName"])]
        grouped[c["complexName"]] = matched
        print(f"{c['complexName']}: {len(matched)}건 ({'~'.join([year_months[-1], year_months[0]])})")

    out_path = OUTPUT_DIR / f"실거래가_전체_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    build_excel(grouped, out_path)
    print(f"\n완료: {out_path}")


if __name__ == "__main__":
    main()
