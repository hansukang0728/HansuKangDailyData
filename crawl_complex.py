"""네이버 부동산(fin.land.naver.com) 단지별 매매/전세 매물 크롤러.

Playwright로 실제 브라우저 컨텍스트에서 article/list API를 호출해 매물을 수집하고,
이전 수집 결과(JSON 스냅샷)와 비교해 신규/삭제 매물을 표시한 엑셀 파일을 만든다.
watchlist.json에 등록된 단지 전체 또는 특정 단지 하나만 골라 수집할 수 있다.
"""
import argparse
import json
import os
import time
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from playwright.sync_api import sync_playwright

from watchlist import load_watchlist

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

# 세션 확보용으로 방문할 지도 페이지 (수원 영통/광교 인근). complexNumber와는 무관하게
# fin.land.naver.com 원본(origin)에서 fetch를 실행하기 위한 용도.
# 주의: /complexes/{번호} 페이지는 financial.pstatic.net 오리진이라 CORS로 막힌다.
MAP_URL = (
    "https://fin.land.naver.com/map?center=3zkve1-2AAmjA&zoom=15.000000000000002"
    "&layer=NobwRAlgJmBcYGMD2BbADgGwKYA8D6UWALgIYQZgA0YaJATiSgM5zjLrY4CSM8AjAAY%2BAJgDsAZjABf"
    "akyz0EACwAK9Ri1jgITAGrkMJOADMSGOdVIAjOGHpEICbNOp1iAVzoA7EpaewidG5YUgC6QA"
)

STEALTH_JS = """
Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
window.chrome = { runtime: {} };
Object.defineProperty(navigator, 'plugins', { get: () => [1, 2, 3, 4, 5] });
Object.defineProperty(navigator, 'languages', { get: () => ['ko-KR', 'ko', 'en-US', 'en'] });
"""

FETCH_JS = """
async (payload) => {
    try {
        const res = await fetch("https://fin.land.naver.com/front-api/v1/complex/article/list", {
            method: "POST",
            headers: {"content-type": "application/json", "accept": "application/json, text/plain, */*"},
            body: JSON.stringify(payload),
            credentials: "include",
        });
        return { status: res.status, text: await res.text() };
    } catch (e) {
        return { status: -1, text: String(e) };
    }
}
"""


@contextmanager
def browser_page():
    """세션이 잡힌 지도 페이지의 Playwright page를 하나 열어 재사용할 수 있게 해준다."""
    with sync_playwright() as p:
        # 클라우드 환경(Claude Code)에는 chromium이 사전 설치돼 있고 playwright가 pip으로
        # 업데이트되면 번들 버전과 어긋나 "Executable doesn't exist"로 실패한다. 그래서
        # 사전 설치된 chromium 실행파일이 있으면 그 경로를 직접 지정한다(CHROMIUM_EXECUTABLE
        # 또는 /opt/pw-browsers/chromium 심볼릭). 로컬(경로 없음)에서는 번들 브라우저로 폴백.
        launch_kwargs = {"headless": False, "args": ["--disable-blink-features=AutomationControlled"]}
        chromium_exe = os.environ.get("CHROMIUM_EXECUTABLE") or "/opt/pw-browsers/chromium"
        if os.path.exists(chromium_exe):
            launch_kwargs["executable_path"] = chromium_exe
        browser = p.chromium.launch(**launch_kwargs)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/150.0.0.0 Safari/537.36"
            ),
            locale="ko-KR",
        )
        context.add_init_script(STEALTH_JS)
        page = context.new_page()
        page.goto(MAP_URL, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(1500)
        try:
            yield page
        finally:
            browser.close()


def fetch_articles(page, complex_number: str, trade_types: list) -> list:
    """이미 열려 있는 page 컨텍스트에서 매물 목록을 페이지네이션 돌며 전부 가져온다."""
    articles = []
    last_info = []
    seed = f"crawl-{int(time.time())}"

    while True:
        payload = {
            "size": 30,
            "complexNumber": complex_number,
            "tradeTypes": trade_types,
            "pyeongTypes": [],
            "dongNumbers": [],
            "userChannelType": "PC",
            "articleSortType": "RANKING_DESC",
            "seed": seed,
            "lastInfo": last_info,
        }
        result = page.evaluate(FETCH_JS, payload)
        if result["status"] != 200:
            raise RuntimeError(f"API 요청 실패: {result['status']} {result['text'][:300]}")

        data = json.loads(result["text"])["result"]
        for item in data["list"]:
            article = item["representativeArticleInfo"]
            dup = item.get("duplicatedArticleInfo")
            if dup:
                article["_realtorCount"] = dup["realtorCount"]
                article["_brokerList"] = [
                    a["brokerInfo"]["brokerageName"] for a in dup["articleInfoList"]
                ]
            else:
                article["_realtorCount"] = 1
                article["_brokerList"] = [article["brokerInfo"]["brokerageName"]]
            articles.append(article)

        if not data.get("hasNextPage"):
            break
        last_info = data["lastInfo"]
        page.wait_for_timeout(1500)  # 페이지 사이 딜레이

    return articles


BROKER_SUFFIXES = [
    "부동산중개법인주식회사", "부동산중개법인", "공인중개사사무소", "부동산중개", "공인중개사", "부동산", "공인",
]


def shorten_broker_name(name: str) -> str:
    for suf in BROKER_SUFFIXES:
        name = name.replace(suf, "")
    return name.strip()


def simplify(article: dict) -> dict:
    space = article["spaceInfo"]
    floor = article["articleDetail"]["floorDetailInfo"]
    building = article["buildingInfo"]
    broker_list = [shorten_broker_name(b) for b in article.get("_brokerList", [article["brokerInfo"]["brokerageName"]])]
    return {
        "articleNumber": article["articleNumber"],
        "complexName": article["complexName"],
        "dong": article["dongName"],
        "floorInfo": article["articleDetail"]["floorInfo"],
        "targetFloor": floor["targetFloor"],
        "totalFloor": floor["totalFloor"],
        "supplySpace": space["supplySpace"],
        "exclusiveSpace": space["exclusiveSpace"],
        "typeName": space.get("nameType", ""),
        "direction": article["articleDetail"]["direction"],
        "dealPrice": article["priceInfo"]["dealPrice"],
        "buildYear": building["buildingConjunctionDate"][:4],
        "elapsedYear": building["approvalElapsedYear"],
        "broker": article["brokerInfo"]["brokerageName"],
        "realtorCount": article.get("_realtorCount", 1),
        "brokerList": broker_list,
        "feature": article["articleDetail"].get("articleFeatureDescription", ""),
        "confirmDate": article["verificationInfo"]["articleConfirmDate"],
    }


def format_price(won: int) -> str:
    eok = won // 100_000_000
    man = (won % 100_000_000) // 10_000
    if man:
        return f"{eok}억 {man:,}만"
    return f"{eok}억"


def format_price_diff(diff: int) -> str:
    if diff == 0:
        return ""
    sign = "+" if diff > 0 else "-"
    eok, man = divmod(abs(diff), 100_000_000)
    man //= 10_000
    parts = [p for p in (f"{eok}억" if eok else "", f"{man:,}만" if man else "") if p]
    return sign + " ".join(parts) if parts else f"{sign}{abs(diff):,}원"


def is_low_floor(target_floor: str) -> bool:
    if target_floor == "저":
        return True
    if target_floor.isdigit():
        return int(target_floor) <= 5
    return False


def save_snapshot(complex_number: str, records: list) -> Path:
    complex_dir = DATA_DIR / complex_number
    complex_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = complex_dir / f"{ts}.json"
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_previous_snapshot(complex_number: str, exclude: Path):
    complex_dir = DATA_DIR / complex_number
    if not complex_dir.exists():
        return None
    files = sorted(f for f in complex_dir.glob("*.json") if f != exclude)
    if not files:
        return None
    return json.loads(files[-1].read_text(encoding="utf-8"))


def assemble_rows(current: list, previous) -> tuple:
    """현재 매물 + (있으면) 사라진 매물에 상태(신규/가격올림/가격내림/변경없음/삭제/최초수집)를 붙여 합친다."""
    if previous is None:
        rows = [dict(r, status="최초수집", priceDiff=0) for r in current]
        return rows, True

    prev_map = {r["articleNumber"]: r for r in previous}
    cur_ids = {r["articleNumber"] for r in current}

    rows = []
    for r in current:
        prev = prev_map.get(r["articleNumber"])
        if prev is None:
            status, diff = "신규", 0
        else:
            diff = r["dealPrice"] - prev["dealPrice"]
            status = "가격올림" if diff > 0 else "가격내림" if diff < 0 else "변경없음"
        rows.append(dict(r, status=status, priceDiff=diff))

    rows += [
        dict(r, status="삭제", priceDiff=0)
        for r in previous
        if r["articleNumber"] not in cur_ids
    ]
    return rows, False


def safe_sheet_name(name: str) -> str:
    for ch in r"[]:*?/\\":
        name = name.replace(ch, "-")
    return name[:31]


THIN_SIDE = Side(style="thin", color="FF000000")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

LEFT_HEADERS = ["번호", "상태", "매물번호", "동", "층"]
RIGHT_HEADERS = [
    "공급면적(㎡)", "전용면적(㎡)", "타입", "방향", "매매가", "가격변동", "중개사무소", "부동산수", "매물확인일",
]
N_LEFT = len(LEFT_HEADERS)
N_RIGHT = len(RIGHT_HEADERS)
N_COLS = N_LEFT + N_RIGHT
ROWS_PER_LISTING = 3  # 기본정보 / 특이사항 / 부동산 목록


def _bordered(ws, row, col, value=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = CELL_BORDER
    return cell


def write_complex_sheet(ws, complex_name: str, rows: list):
    # 삭제된 매물은 맨 아래로 몰아서 배치, 그 안에서도 가격순
    rows_sorted = sorted(rows, key=lambda r: (r["status"] == "삭제", r["dealPrice"]))
    first = rows_sorted[0] if rows_sorted else None

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_COLS)
    title_cell = ws.cell(row=1, column=1, value=complex_name)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="left")

    info_text = (
        f"준공년도 {first['buildYear']}년 · 경과 {first['elapsedYear']}년 · 총 {len(rows_sorted)}건"
        if first is not None
        else "매물 없음"
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N_COLS)
    ws.cell(row=2, column=1, value=info_text).font = Font(italic=True)

    header_row = 4
    for col, name in enumerate(LEFT_HEADERS + RIGHT_HEADERS, start=1):
        cell = _bordered(ws, header_row, col, name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"

    red_bold = Font(color="FFFF0000", bold=True)
    blue_bold = Font(color="FF0000FF", bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for offset, r in enumerate(rows_sorted):
        top_row = header_row + 1 + offset * ROWS_PER_LISTING
        feature_row = top_row + 1
        broker_row = top_row + 2
        bottom_row = broker_row

        # 좌측: 번호/상태/매물번호/동/층 — 세 줄(기본정보+특이사항+부동산목록)에 걸쳐 세로 병합
        left_values = [offset + 1, r["status"], r["articleNumber"], r["dong"], r["floorInfo"]]
        for col, value in enumerate(left_values, start=1):
            ws.merge_cells(start_row=top_row, start_column=col, end_row=bottom_row, end_column=col)
            cell = _bordered(ws, top_row, col, value)
            cell.alignment = center
            for extra_row in (feature_row, broker_row):
                _bordered(ws, extra_row, col)  # 병합 범위 나머지 줄도 테두리 필요
            if col == 5 and is_low_floor(r["targetFloor"]):
                cell.font = red_bold

        # 우측 첫줄: 면적/타입/방향/매매가/가격변동/중개사무소/부동산수/매물확인일
        # realtorCount/brokerList는 구버전 스냅샷(삭제 상태로 남아있는 옛 데이터)에는
        # 없을 수 있어 기본값으로 보완한다.
        right_values = [
            r["supplySpace"], r["exclusiveSpace"], r["typeName"], r["direction"],
            format_price(r["dealPrice"]), format_price_diff(r["priceDiff"]),
            r["broker"], r.get("realtorCount", 1), r["confirmDate"],
        ]
        price_offset = 4  # right_values 중 매매가 위치
        for i, value in enumerate(right_values):
            cell = _bordered(ws, top_row, N_LEFT + 1 + i, value)
            if i == price_offset:
                cell.font = blue_bold

        # 우측 둘째줄: 특이사항 — 가로 병합 한 줄
        ws.merge_cells(start_row=feature_row, start_column=N_LEFT + 1, end_row=feature_row, end_column=N_COLS)
        feat_cell = _bordered(ws, feature_row, N_LEFT + 1, r["feature"])
        feat_cell.alignment = Alignment(horizontal="left")
        for col in range(N_LEFT + 2, N_COLS + 1):
            _bordered(ws, feature_row, col)

        # 우측 셋째줄: 부동산 목록(축약명, 쉼표 구분) — 가로 병합 한 줄
        ws.merge_cells(start_row=broker_row, start_column=N_LEFT + 1, end_row=broker_row, end_column=N_COLS)
        broker_cell = _bordered(ws, broker_row, N_LEFT + 1, ", ".join(r.get("brokerList", [r["broker"]])))
        broker_cell.alignment = Alignment(horizontal="left")
        broker_cell.font = Font(size=9, italic=True)
        for col in range(N_LEFT + 2, N_COLS + 1):
            _bordered(ws, broker_row, col)


def build_excel(complex_name: str, rows: list, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "매물목록"
    write_complex_sheet(ws, complex_name, rows)
    wb.save(out_path)


def build_combined_excel(complex_results: list, out_path: Path):
    """[(complex_name, rows), ...] 를 단지별 시트로 묶은 통합 워크북을 만든다."""
    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()
    for complex_name, rows in complex_results:
        sheet_name = safe_sheet_name(complex_name)
        base_name = sheet_name
        suffix = 2
        while sheet_name in used_names:
            sheet_name = safe_sheet_name(f"{base_name}({suffix})")
            suffix += 1
        used_names.add(sheet_name)
        ws = wb.create_sheet(title=sheet_name)
        write_complex_sheet(ws, complex_name, rows)
    wb.save(out_path)


def crawl_one(page, complex_number: str, complex_name: str, region: str, trade_types: list) -> dict:
    print(f"[수집] {complex_name}({complex_number}) ...")
    raw_articles = fetch_articles(page, complex_number, trade_types)
    print(f"  -> {len(raw_articles)}건 수집")

    records = [simplify(a) for a in raw_articles]

    snapshot_path = save_snapshot(complex_number, records)
    previous = load_previous_snapshot(complex_number, exclude=snapshot_path)
    rows, is_first = assemble_rows(records, previous)
    new_count = sum(1 for r in rows if r["status"] == "신규")
    removed_count = sum(1 for r in rows if r["status"] == "삭제")
    up_count = sum(1 for r in rows if r["status"] == "가격올림")
    down_count = sum(1 for r in rows if r["status"] == "가격내림")
    if is_first:
        print("  -> 최초 수집이라 비교 대상 없음")
    else:
        print(f"  -> 신규 {new_count}건, 삭제 {removed_count}건, 가격올림 {up_count}건, 가격내림 {down_count}건")

    complex_dir = OUTPUT_DIR / f"{region}_{complex_number}_{complex_name}"
    complex_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = complex_dir / f"{complex_name}_{ts}.xlsx"
    build_excel(complex_name, rows, out_path)
    print(f"  -> 저장: {out_path}")

    return {
        "complex_number": complex_number,
        "complex_name": complex_name,
        "count": len(records),
        "new": new_count,
        "removed": removed_count,
        "path": out_path,
        "rows": rows,
    }


def main():
    parser = argparse.ArgumentParser(description="네이버 부동산 단지별 매물 크롤러")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--complex-number", help="단지 번호로 1회성 수집 (관심단지 등록 여부 무관)")
    group.add_argument("--name", help="관심단지 목록에서 이름으로 찾아 수집")
    group.add_argument("--all", action="store_true", help="관심단지 전체 수집")
    parser.add_argument("--complex-name", default=None, help="--complex-number와 함께 사용, 결과 파일명에 쓸 이름")
    parser.add_argument("--region", default="기타", help="--complex-number와 함께 사용, 결과 파일명에 쓸 지역명")
    parser.add_argument("--trade-types", nargs="+", default=["A1"], help="--complex-number와 함께 사용")
    args = parser.parse_args()

    DATA_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)

    if args.all:
        watchlist = load_watchlist()
        if not watchlist:
            print("관심단지가 비어 있습니다. watchlist.py add 로 먼저 등록하세요.")
            return
        targets = [
            (c["complexNumber"], c["complexName"], c.get("region", "기타"), c.get("tradeTypes", ["A1"]))
            for c in watchlist
        ]
    elif args.name:
        watchlist = load_watchlist()
        match = next((c for c in watchlist if c["complexName"] == args.name), None)
        if not match:
            print(f"관심단지에서 '{args.name}'을 찾을 수 없습니다. watchlist.py list 로 확인하세요.")
            return
        targets = [(
            match["complexNumber"], match["complexName"],
            match.get("region", "기타"), match.get("tradeTypes", ["A1"]),
        )]
    else:
        targets = [(args.complex_number, args.complex_name or args.complex_number, args.region, args.trade_types)]

    results = []
    with browser_page() as page:
        for i, (num, name, region, trade_types) in enumerate(targets):
            results.append(crawl_one(page, num, name, region, trade_types))
            if i < len(targets) - 1:
                page.wait_for_timeout(3000)  # 단지 사이 딜레이

    if len(results) > 1:
        print("\n=== 요약 ===")
        for r in results:
            print(f"{r['complex_name']}: {r['count']}건 (신규 {r['new']}, 삭제 {r['removed']}) -> {r['path'].relative_to(OUTPUT_DIR)}")

        combined_path = OUTPUT_DIR / f"전체_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        build_combined_excel([(r["complex_name"], r["rows"]) for r in results], combined_path)
        print(f"통합 파일: {combined_path}")


if __name__ == "__main__":
    main()
